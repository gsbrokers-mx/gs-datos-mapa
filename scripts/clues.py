"""CLUES (Secretaría de Salud): hospitales y clínicas, distinguiendo los institucionales."""
import re
from functools import lru_cache

import openpyxl
import requests

from comun import UA, descargar, log, norm, punto, titulo

FUENTE = "CLUES"
PAGINA = "http://www.dgis.salud.gob.mx/contenidos/intercambio/clues_gobmx.html"

# Clave de institución -> nombre corto. SMP (servicios médicos privados) queda como privado.
INSTITUCIONES = {
    "IMS": "IMSS", "IMB": "IMSS-Bienestar", "IST": "ISSSTE", "SSA": "Secretaría de Salud",
    "SME": "Servicios Médicos Estatales", "SMM": "Servicios Médicos Municipales", "DIF": "DIF",
    "CRO": "Cruz Roja", "CIJ": "Centros de Integración Juvenil", "HUN": "Hospital Universitario",
    "PMX": "Pemex", "SDN": "Sedena", "SMA": "Semar",
}
# Tipologías que no sirven como referencia en el mapa
EXCLUIR_TIPOLOGIA = re.compile(r"ADYACENTE A FARMACIA|UNIDAD MOVIL|CENTRO DE TRABAJO|ALMACEN|OFICINAS")
RAZON_SOCIAL = re.compile(r",?\s+(S\.?\s?A\.?\s?P?\.?\s?I?\.?\s?(DE\s+C\.?\s?V\.?)?|S\.?\s?C\.?|A\.?\s?C\.?|S\.?\s?DE\s+R\.?\s?L\.?.*)\s*$")


@lru_cache(maxsize=1)
def _archivo():
    html = requests.get(PAGINA, headers=UA, timeout=60).text
    m = re.search(r'https?://[^"\']+ESTABLECIMIENTO_SALUD_\d{6}\.xlsx', html)
    if not m:
        raise RuntimeError("No encontré el enlace del catálogo CLUES en la página de la DGIS")
    log(f"   CLUES: {m.group(0).rsplit('/', 1)[-1]}")
    return descargar(m.group(0), "clues.xlsx", max_horas=20)


def salud(zona):
    wb = openpyxl.load_workbook(_archivo(), read_only=True)
    hoja = next(n for n in wb.sheetnames if n.startswith("CLUES"))
    filas = wb[hoja].iter_rows(values_only=True)
    enc = next(filas)
    puntos = []
    for fila in filas:
        r = dict(zip(enc, fila))
        if str(r["CLAVE DE LA ENTIDAD"]).zfill(2) != zona["entidad"]:
            continue
        if str(r["CLAVE DEL MUNICIPIO"]).zfill(3) not in zona["municipios"]:
            continue
        if norm(r["ESTATUS DE OPERACION"]) != "EN OPERACION":
            continue
        try:
            lat, lng = float(r["LATITUD"]), float(r["LONGITUD"])
        except (TypeError, ValueError):
            continue
        tipo = norm(r["NOMBRE TIPO ESTABLECIMIENTO"])
        tipologia = norm(r["NOMBRE DE TIPOLOGIA"])
        if tipo in ("DE APOYO", "DE ASISTENCIA SOCIAL") or EXCLUIR_TIPOLOGIA.search(tipologia):
            continue
        hospital = tipo == "DE HOSPITALIZACION"
        inst = INSTITUCIONES.get(r["CLAVE DE LA INSTITUCION"])
        nombre = (r["NOMBRE COMERCIAL"] or "").strip() or (r["NOMBRE DE LA UNIDAD"] or "").strip()
        nombre = titulo(RAZON_SOCIAL.sub("", nombre.upper()).strip(" ,."))
        if inst and inst not in nombre and not re.search(r"\b(IMSS|ISSSTE|CRUZ|DIF|UMF)\b", norm(nombre)):
            nombre = f"{nombre} ({inst})" if len(nombre) < 60 else nombre
        if re.search(r"DENTAL|ODONTOL", tipologia + " " + norm(nombre)):
            clave = "salud:dentales"
        else:
            clave = "salud:hospitales" if hospital else "salud:clinicas"
        puntos.append(punto(nombre, lat, lng, clave, hospital or bool(inst), FUENTE,
                            institucion=inst or "Privado"))
    return puntos
